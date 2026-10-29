"""Make the manufacturer Excel file: the labels (columns) the importer will read, with drop-downs and two example
rows. Fill it in, then upload it to the server.

    uv run python scripts/make_manufacturer_sheet.py                      # writes data/manufacturers.xlsx
    uv run python scripts/make_manufacturer_sheet.py path/to/file.xlsx

Rows whose id starts with "example-" are samples and are ignored by the importer.
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# (column id, width, help, allowed values or None, example 1, example 2)
FACTORIES = [
    ("id", 22, "Short unique name, lowercase, no spaces (used everywhere). Never change it later.", None,
     "example-factory-a", "example-factory-b"),
    ("name", 24, "Factory name in English", None, "Guangzhou Sunrise Garments", "Yiwu Bright Packaging"),
    ("name_zh", 20, "Factory name in Chinese (as on their business card)", None, "广州朝阳服饰有限公司", "义乌明亮包装厂"),
    ("active", 9, "no = never contacted", ["yes", "no"], "yes", "yes"),
    ("channel", 14, "How Orion reaches them. wecom = WeChat through WeCom; manual_relay = a team member copies the message",
     ["wecom", "email", "manual_relay"], "wecom", "email"),
    ("contact", 26, "email: their email address. wecom: leave empty (they connect through our WeChat link). "
                    "manual_relay: their WeChat id or phone", None, "", "sales@example.cn"),
    ("contact_person", 18, "Who we talk to", None, "王经理", "Lily Chen"),
    ("contact_person_title", 14, "Their role (经理 / 老板 / 业务员 / 跟单)", None, "经理", "业务员"),
    ("language", 10, "Language of the messages we send them", ["zh", "en"], "zh", "zh"),
    ("city", 14, "City or province", None, "Guangzhou", "Yiwu"),
    ("website_or_1688", 30, "Their website, 1688 or Alibaba page", None, "https://example.1688.com", ""),
    ("methods", 30, "What they can do, comma separated (embroidery, screen print, DTG, woven label...)", None,
     "embroidery, screen print, woven label", "offset print, die cut"),
    ("moq", 10, "Their minimum order quantity (a fact about them, not a price)", None, 100, 500),
    ("lead_days", 10, "Normal production time in days", None, 20, 15),
    ("sample_days", 10, "Normal sample time in days", None, 7, 5),
    ("notes_internal", 40, "For the team only. Not shown to the AI or to the factory.", None,
     "Good quality, slow on holidays", ""),
]

PRODUCTS = [
    ("manufacturer_id", 22, "The factory's id from the Factories sheet", None, "example-factory-a", "example-factory-a"),
    ("market", 14, "apparel, ecom_packaging or tech_accessories", None, "apparel", "apparel"),
    ("option", 16, "Product, e.g. hoodie, tee, polo, abaya, tote. Empty = every product in the market", None,
     "hoodie", ""),
    ("rank", 8, "1 = contacted first. Lower numbers first", None, 1, 2),
    ("notes", 40, "Why this rank (team only)", None, "Best hoodie quality", "Backup for any apparel"),
]

# How the agent talks to this factory. No price, discount or negotiation settings exist on purpose (rule 1).
PROFILE = [
    ("manufacturer_id", 22, "The factory's id from the Factories sheet", None, "example-factory-a", "example-factory-b"),
    ("tone", 12, "formal = polite business Chinese; friendly = warm and relaxed; direct = short and to the point",
     ["formal", "friendly", "direct"], "formal", "friendly"),
    ("address_as", 16, "How the agent addresses them", None, "王经理", "Lily"),
    ("greeting", 30, "Opening line (leave empty for the default)", None, "王经理您好，", "Hi Lily!"),
    ("sign_off", 30, "Closing line (leave empty for the default)", None, "谢谢！Orion 团队", "Thanks! Orion team"),
    ("introduce_orion", 16, "When to explain who Orion is", ["first_message", "always", "never"],
     "first_message", "first_message"),
    ("message_length", 14, "short = few lines; normal; detailed = full explanations", ["short", "normal", "detailed"],
     "normal", "short"),
    ("emoji", 8, "May the agent use emoji", ["no", "few"], "no", "few"),
    ("one_topic_per_message", 12, "yes = never mix topics in one message (good for the 5-message WeChat limit)",
     ["yes", "no"], "yes", "no"),
    ("working_hours_china", 16, "When we send, China time (24h). Outside this, messages wait", None, "09:00-18:00", "09:00-20:00"),
    ("reply_wait_hours", 12, "Hours before the first reminder (default 72 = 3 days)", None, 72, 48),
    ("reminder_style", 14, "gentle = polite nudge; firm = clear that we need an answer", ["gentle", "firm"],
     "gentle", "gentle"),
    ("always_include", 40, "Things every message to them must contain", None,
     "Order number at the start. Attach the placement drawing.", ""),
    ("never_say", 40, "Words or topics to avoid. (Prices, discounts and payment are always blocked by code.)", None,
     "Don't mention other factories", ""),
    ("special_instructions", 60, "Anything else about how to talk to this factory, in plain words (English or Chinese)",
     None, "Boss reads WeChat only after lunch. Always ask for a photo of the finished sample.", ""),
    ("sample_message", 60, "A real message that worked well with them (Chinese). The agent copies its style", None,
     "王经理您好，我们有一个新的卫衣订单，规格和图纸见附件，请确认是否可以生产，谢谢！", ""),
]

READ_ME = [
    ("Manufacturer sheet", ""),
    ("", ""),
    ("Sheets", "Factories = who they are and how to reach them. Products = which products each factory makes and "
               "the order we contact them in. Agent profile = how the AI should talk to each one."),
    ("One row per", "Factories: one row per factory. Products: one row per factory and product. "
                    "Agent profile: one row per factory."),
    ("id", "Each factory has one id. Use the same id on all three sheets. Do not change it after the first upload."),
    ("Examples", "Rows whose id starts with example- are samples. Delete them or leave them: they are ignored."),
    ("Empty cells", "Optional columns can stay empty; the default is used."),
    ("Prices", "This file has no prices on purpose. The AI never talks about prices, discounts or payment. "
               "It only records the prices factories send."),
    ("WeChat", "For channel = wecom the factory first opens our WeChat customer-service link. Give each factory "
               "its own link (the link carries its id), and the connection is made when they open it."),
    ("Uploading", "Put the file on the server and run: uv run python -m orion.relay import manufacturers.xlsx "
                  "(it shows what would change first and saves nothing until you confirm)."),
]

HEAD = PatternFill("solid", fgColor="1F2A44")
HELP = PatternFill("solid", fgColor="E8ECF5")
SAMPLE = PatternFill("solid", fgColor="FFF8E1")


def sheet(wb: Workbook, title: str, columns: list, first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for c, (key, width, help_text, allowed, *examples) in enumerate(columns, 1):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = max(width, 12)
        head = ws.cell(1, c, key)
        head.font, head.fill = Font(bold=True, color="FFFFFF"), HEAD
        head.alignment = Alignment(vertical="center")
        hint = ws.cell(2, c, help_text)
        hint.font, hint.fill = Font(italic=True, size=9, color="444444"), HELP
        hint.alignment = Alignment(wrap_text=True, vertical="top")
        for r, value in enumerate(examples, 3):
            cell = ws.cell(r, c, value if value != "" else None)
            cell.fill, cell.alignment = SAMPLE, Alignment(wrap_text=True, vertical="top")
        if allowed:
            check = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
            ws.add_data_validation(check)
            check.add(f"{letter}3:{letter}500")
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 64
    ws.freeze_panes = "B3"
    return ws


def main() -> None:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "data/manufacturers.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    sheet(wb, "Factories", FACTORIES, first=True)
    sheet(wb, "Products", PRODUCTS)
    sheet(wb, "Agent profile", PROFILE)
    info = wb.create_sheet("Read me", 0)
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 110
    for r, (a, b) in enumerate(READ_ME, 1):
        info.cell(r, 1, a).font = Font(bold=True, size=14 if r == 1 else 11)
        info.cell(r, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
